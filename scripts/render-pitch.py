#!/usr/bin/env python3
"""Render docs/pitch/deck.html to PDF, slide PNGs and a PPTX with speaker notes.

Needs Python 3, Pillow, python-pptx, Playwright for Python and Microsoft Edge or
Chrome (`--channel`). Fonts: HarmonyOS Sans SC if installed, otherwise Segoe UI
or the system sans-serif.

  python scripts/render-pitch.py                 # deck -> dist/ and docs/pitch/
  python scripts/render-pitch.py --screens       # first refresh docs/pitch/img from the gallery

--screens copies the phone screenshots the deck uses from the full-resolution
gallery (`.cache/gallery-raw`, written by scripts/capture-gallery.py) and the
demo stills (`.cache/demo-video/stills`, written by scripts/demo-video/drive.py).

Outputs:
  dist/SafeMesh-pitch.pdf, dist/SafeMesh-pitch.pptx, dist/pitch/slide-NN.png
  docs/pitch/SafeMesh-pitch.pdf and docs/pitch/slides/NN.jpg (committed previews)
"""

import argparse
import html
import re
from pathlib import Path
import shutil

from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
PITCH = ROOT / 'docs' / 'pitch'
SCREEN_HEIGHT = 2856
# deck image <- (gallery variant, screen) or a demo still
SCREENS = {
    'screen-message.jpg': ('en-light', '17-message-detail'),
    'screen-map.jpg': ('en-light', '02-map'),
    'screen-point.jpg': ('en-light', '05-point-detail'),
    'screen-home-en-light.jpg': ('en-light', '16-home-inbox'),
    'screen-home-pl-dark.jpg': ('pl-dark', '16-home-inbox'),
}
STILLS = {'demo-a.jpg': 'a', 'demo-b.jpg': 'b', 'demo-c.jpg': 'c'}


def phone_image(source, target, width=660):
    with Image.open(source) as image:
        image = image.convert('RGB')
        image = image.crop((0, 0, image.width, min(image.height, round(image.width * SCREEN_HEIGHT / 1320))))
        image = image.resize((width, round(image.height * width / image.width)), Image.LANCZOS)
        image.save(target, quality=90, optimize=True, progressive=True)


def refresh_screens(gallery, stills):
    for name, (variant, screen) in SCREENS.items():
        source = gallery / variant / f'{screen}.png'
        if not source.exists():
            raise SystemExit(f'missing {source}; run scripts/capture-gallery.py first')
        phone_image(source, PITCH / 'img' / name)
        print('screen', name, '<-', source.relative_to(ROOT))
    for name, node in STILLS.items():
        source = stills / f'{node}.png'
        if source.exists():
            phone_image(source, PITCH / 'img' / name)
            print('still', name, '<-', source.relative_to(ROOT))
        else:
            print('still', name, 'kept (no', source.relative_to(ROOT), ')')


def notes_and_titles(deck_html):
    sections = re.findall(r'<section class="slide[^"]*" data-title="([^"]*)">(.*?)</section>', deck_html, re.S)
    result = []
    for title, body in sections:
        match = re.search(r'<aside class="notes">(.*?)</aside>', body, re.S)
        text = re.sub(r'<[^>]+>', '', match.group(1)) if match else ''
        result.append((html.unescape(title), html.unescape(' '.join(text.split()))))
    return result


def render(channel):
    from playwright.sync_api import sync_playwright
    out = ROOT / 'dist' / 'pitch'
    out.mkdir(parents=True, exist_ok=True)
    (PITCH / 'slides').mkdir(exist_ok=True)
    url = (PITCH / 'deck.html').resolve().as_uri() + '?render'
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel=channel)
        page = browser.new_page(viewport={'width': 1920, 'height': 1080}, device_scale_factor=1)
        page.goto(url)
        page.wait_for_load_state('networkidle')
        page.evaluate('document.fonts.ready')
        page.wait_for_timeout(500)
        slides = page.query_selector_all('section.slide')
        for number, slide in enumerate(slides, 1):
            png = out / f'slide-{number:02d}.png'
            slide.screenshot(path=str(png))
            with Image.open(png) as image:
                image.convert('RGB').resize((960, 540), Image.LANCZOS).save(
                    PITCH / 'slides' / f'{number:02d}.jpg', quality=86, optimize=True, progressive=True)
        page.emulate_media(media='print')
        page.pdf(path=str(ROOT / 'dist' / 'SafeMesh-pitch.pdf'), width='1920px', height='1080px',
                 print_background=True, prefer_css_page_size=True)
        browser.close()
    shutil.copyfile(ROOT / 'dist' / 'SafeMesh-pitch.pdf', PITCH / 'SafeMesh-pitch.pdf')
    return len(slides)


def build_pptx(count, notes):
    from pptx import Presentation
    from pptx.util import Emu
    deck = Presentation()
    deck.slide_width, deck.slide_height = Emu(12192000), Emu(6858000)
    blank = deck.slide_layouts[6]
    for number in range(1, count + 1):
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(ROOT / 'dist' / 'pitch' / f'slide-{number:02d}.png'), 0, 0,
                                 width=deck.slide_width, height=deck.slide_height)
        title, text = notes[number - 1]
        slide.notes_slide.notes_text_frame.text = text
        slide.name = title
    deck.core_properties.title = 'SafeMesh pitch deck'
    deck.core_properties.subject = 'HackYeah 2026 · Huawei Imagine What’s Next'
    deck.save(ROOT / 'dist' / 'SafeMesh-pitch.pptx')


def write_readme(count, notes):
    lines = ['# SafeMesh pitch deck', '',
             f'{count} slides for the HackYeah 2026 jury. **[Open the PDF](SafeMesh-pitch.pdf)**, or open '
             '[`deck.html`](deck.html) in a browser (arrow keys, F for full screen). The PPTX with speaker notes '
             'and the PDF are also attached to the latest GitHub release.', '',
             'Rendered by `python scripts/render-pitch.py`. Phone images are real emulator screenshots from '
             '[the screen gallery](../gallery/README.md) and the demo recording, including the three phones on the title slide.', '']
    for number in range(1, count + 1):
        title, text = notes[number - 1]
        lines += [f'### {number}. {title}', '', f'![Slide {number}: {title}](slides/{number:02d}.jpg)', '',
                  f'> {text}', '']
    (PITCH / 'README.md').write_text('\n'.join(lines), encoding='utf-8', newline='\n')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--screens', action='store_true', help='Refresh docs/pitch/img from the gallery and stills.')
    parser.add_argument('--gallery', type=Path, default=ROOT / '.cache' / 'gallery-raw')
    parser.add_argument('--stills', type=Path, default=ROOT / '.cache' / 'demo-video' / 'stills')
    parser.add_argument('--channel', default='msedge', help='Playwright browser channel: msedge or chrome.')
    args = parser.parse_args()
    if args.screens:
        refresh_screens(args.gallery, args.stills)
    notes = notes_and_titles((PITCH / 'deck.html').read_text(encoding='utf-8'))
    count = render(args.channel)
    if count != len(notes):
        raise SystemExit(f'{count} slides rendered but {len(notes)} have notes')
    build_pptx(count, notes)
    write_readme(count, notes)
    print(f'{count} slides -> dist/SafeMesh-pitch.pdf, dist/SafeMesh-pitch.pptx, docs/pitch/')


if __name__ == '__main__':
    main()
